from pathlib import Path
import pytest
from analyzer.parser import parse
from analyzer.models import Analysis
from analyzer.validation import validate
ROOT=Path(__file__).resolve().parents[1]
def sample():
    return parse(ROOT/'examples/coupon_srs.txt','SRS-1','1.0'),Analysis.model_validate_json((ROOT/'examples/demo_analysis.json').read_text())
def test_sample_valid():
    d,a=sample();assert validate(a,d)==[]
def test_fake_quote_rejected():
    d,a=sample();a.rules[0].source_refs[0].quote='Invented rule'
    with pytest.raises(ValueError,match='Quote'):validate(a,d)
def test_orphan_scenario_rejected():
    d,a=sample();a.scenarios[0].rule_ids=['MISSING']
    with pytest.raises(ValueError,match='rule không tồn tại'):validate(a,d)
def test_unknown_requires_issue():
    d,a=sample();a.issues=[]
    with pytest.raises(ValueError,match='issue tương ứng'):validate(a,d)
def test_html_escapes_model_content(tmp_path):
    from analyzer.report import write_report
    d,a=sample();a.summary='<script>alert(1)</script>'
    write_report({'provider':'TEST','model':'none','analysis':a.model_dump(),'source_document':d.model_dump(),'validation_warnings':[]},tmp_path)
    assert '<script>' not in (tmp_path/'report.html').read_text()
def test_html_matches_demo_report_sections(tmp_path):
    from analyzer.report import write_report
    d,a=sample()
    write_report({'provider':'GEMINI','model':'test-model','analysis':a.model_dump(),'source_document':d.model_dump(),'validation_warnings':[],'pipeline_stats':{'chunks':1}},tmp_path)
    report=(tmp_path/'report.html').read_text(encoding='utf-8')
    assert '<p>Provider: GEMINI | Model: test-model | REVIEW_REQUIRED</p>' in report
    assert '<h2>Business rules (3)</h2>' in report
    assert '<h2>Test scenarios (4)</h2>' in report
    assert '<h2>Issues cần xác nhận (2)</h2>' in report
    assert '<h2>Source blocks</h2>' in report
    assert '<h2>Validation</h2>' in report
    assert '<h2>Pipeline</h2>' not in report
    assert '<h2>Module ' not in report
def test_docx_preserves_table_order(tmp_path):
    from docx import Document
    d=Document();d.add_paragraph('First requirement');d.add_table(rows=1,cols=1).cell(0,0).text='Table rule';d.add_paragraph('Last requirement');f=tmp_path/'srs.docx';d.save(f)
    blocks=parse(f,'DOC','1').blocks
    assert [b.text for b in blocks]==['First requirement','Table rule','Last requirement']
def test_gemini_adapter_with_mock(monkeypatch):
    from analyzer import gemini
    from google.genai import types
    d,a=sample();calls=[]
    class Models:
        def generate_content(self,**kwargs):
            calls.append(kwargs)
            return types.GenerateContentResponse(candidates=[types.Candidate(content=types.Content(parts=[types.Part(text=a.model_dump_json())],role='model'))])
    class FakeClient:
        def __init__(self,**kwargs):self.models=Models()
        def __enter__(self):return self
        def __exit__(self,*args):pass
    monkeypatch.setattr(gemini.genai,'Client',FakeClient)
    result=gemini.generate(d,'fake-key-for-unit-test','test-model')
    assert result.rules[0].rule_id=='R1'
    assert calls[0]['config'].response_mime_type=='application/json'
    schema=calls[0]['config'].response_schema
    assert schema['properties']['rules']['items']['$ref']=='#/$defs/Rule'
    def assert_no_additional_properties(value):
        if isinstance(value,dict):
            assert 'additionalProperties' not in value
            for child in value.values():assert_no_additional_properties(child)
        elif isinstance(value,list):
            for child in value:assert_no_additional_properties(child)
    assert_no_additional_properties(schema)

def test_missing_key_clear_error():
    from analyzer.gemini import generate
    d,_=sample()
    with pytest.raises(ValueError,match='GEMINI_API_KEY'):generate(d,'','test-model')
def test_real_srs_chunk_coverage():
    from analyzer.chunking import split_document
    d=parse(ROOT/'examples/SRS_Website_ban_truyen_tranh_online_UC.docx','SRS-COMIC','1.0')
    chunks=split_document(d)
    ids=[b.block_id for c in chunks for b in c['blocks']]
    assert len(ids)==len(set(ids))==len(d.blocks)
    assert sum(len(b.text) for b in d.blocks)>40000
    assert len([b for b in d.blocks if b.use_case])==41
    assert any(b.module=='M07' and b.use_case=='UC-14' for b in d.blocks)
    assert all(sum(len(b.text) for b in c['blocks'])<=12000 for c in chunks)
def test_merge_duplicate_local_ids():
    from analyzer.merge import merge_analyses
    d,a=sample();out=merge_analyses([a,a])
    assert len({r.rule_id for r in out.rules})==6
    assert validate(out,d)==[]
def test_pipeline_mock_and_resume(tmp_path):
    from analyzer.pipeline import analyze_full
    d,a=sample();calls=[]
    def generate(doc,key,model,targets):calls.append(targets);return a
    result,warnings,stats=analyze_full(d,'fake','mock',tmp_path,skip_cross=True,generator=generate)
    assert stats['submitted_target_blocks']==len(d.blocks)
    analyze_full(d,'fake','mock',tmp_path,resume=True,skip_cross=True,generator=generate)
    assert len(calls)==1
def test_pipeline_retries_validation_with_feedback(tmp_path):
    from analyzer.pipeline import analyze_full
    d,a=sample();calls=[]
    def generate(doc,key,model,targets,validation_feedback=None):
        calls.append(validation_feedback)
        result=a.model_copy(deep=True)
        if len(calls)==1:result.rules[0].source_refs[0].quote='Not an exact quote'
        return result
    analyze_full(d,'fake','mock',tmp_path,skip_cross=True,generator=generate)
    assert len(calls)==2
    assert calls[0] is None
    assert 'Quote không nguyên văn' in calls[1]
    assert 'REQ-COUPON-01' in calls[1]
    assert 'Not an exact quote' in calls[1]
